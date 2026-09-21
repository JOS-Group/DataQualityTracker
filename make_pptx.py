from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

WHITE = RGBColor(255, 255, 255)
OFFWHITE = RGBColor(251, 248, 240)
GOLD1 = RGBColor(245, 215, 122)
GOLD2 = RGBColor(215, 170, 56)
GOLD3 = RGBColor(167, 121, 24)
GOLD4 = RGBColor(111, 77, 0)
INK = RGBColor(28, 26, 23)
MUTED = RGBColor(95, 89, 79)
LINE = RGBColor(238, 221, 176)


def add_textbox(slide, left, top, width, height, text, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    return box


def add_header(slide, title, tag):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.45), prs.slide_height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = GOLD2
    bar.line.fill.background()

    add_textbox(slide, Inches(0.8), Inches(0.45), Inches(2.2), Inches(0.4), 'DART', 20, GOLD4, True)
    add_textbox(slide, Inches(10.7), Inches(0.52), Inches(2.0), Inches(0.3), tag.upper(), 10, GOLD4, True, PP_ALIGN.RIGHT)
    add_textbox(slide, Inches(0.8), Inches(1.0), Inches(9.8), Inches(0.8), title, 25, INK, True)


def add_card(slide, left, top, width, height, title, body):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid(); shape.fill.fore_color.rgb = OFFWHITE
    shape.line.color.rgb = LINE
    shape.line.width = Pt(1.2)
    add_textbox(slide, left + Inches(0.2), top + Inches(0.16), width - Inches(0.35), Inches(0.35), title, 13, GOLD4, True)
    add_textbox(slide, left + Inches(0.2), top + Inches(0.56), width - Inches(0.35), Inches(0.55), body, 10, MUTED)


def add_metric(slide, left, top, width, height, num, label):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid(); shape.fill.fore_color.rgb = OFFWHITE
    shape.line.color.rgb = LINE
    add_textbox(slide, left + Inches(0.15), top + Inches(0.1), width - Inches(0.2), Inches(0.42), num, 28, GOLD4, True)
    add_textbox(slide, left + Inches(0.15), top + Inches(0.56), width - Inches(0.2), Inches(0.28), label, 10, MUTED)

# Slide 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = WHITE
add_header(slide, 'DART keeps data quality operational, measurable, and fixable.', 'product overview')
add_textbox(slide, Inches(0.9), Inches(2.8), Inches(7.1), Inches(1.2), 'A data assurance and reconciliation platform that helps teams detect issues, assign ownership, prioritize impact, and turn operational data into actionable remediation plans.', 18, MUTED)

card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.0), Inches(1.3), Inches(3.6), Inches(4.2))
card.fill.solid(); card.fill.fore_color.rgb = OFFWHITE
card.line.color.rgb = LINE
card.line.width = Pt(1.2)
add_textbox(slide, Inches(9.35), Inches(1.6), Inches(2.8), Inches(0.5), 'EXECUTIVE SNAPSHOT', 9, GOLD4, True)
add_textbox(slide, Inches(9.35), Inches(2.3), Inches(2.7), Inches(0.8), '93%', 28, GOLD4, True)
add_textbox(slide, Inches(9.35), Inches(3.1), Inches(2.9), Inches(0.5), 'quality visibility', 14, MUTED)
for idx, item in enumerate([('Issue detection', 'Live'), ('Impact scoring', 'Included'), ('AI summaries', 'Enabled'), ('Email automation', 'Built in')]):
    y = Inches(4.0 + idx * 0.52)
    row = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.2), y, Inches(3.0), Inches(0.34))
    row.fill.solid(); row.fill.fore_color.rgb = WHITE; row.line.fill.background()
    add_textbox(slide, Inches(9.3), y - Inches(0.04), Inches(1.9), Inches(0.3), item[0], 10, MUTED)
    add_textbox(slide, Inches(11.1), y - Inches(0.04), Inches(1.0), Inches(0.3), item[1], 10, GOLD4, True, PP_ALIGN.RIGHT)

for i, (num, label) in enumerate([('3','core workspaces'), ('50','state comparisons'), ('9','issue categories')]):
    x = Inches(0.9 + i * 2.75)
    add_metric(slide, x, Inches(5.4), Inches(2.3), Inches(1.2), num, label)

# Slide 2
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = WHITE
add_header(slide, 'Most organizations lose time before they lose trust.', 'the problem')
add_textbox(slide, Inches(0.8), Inches(2.0), Inches(8.6), Inches(1.0), 'Data quality issues create compounding drag across reconciliations, reporting, payments, compliance, and executive confidence.', 18, MUTED)
items = [
    ('Broken reconciliations', 'Claims, ledgers, and source records do not always line up, creating pressure in reporting and finance workflows.'),
    ('Unclear ownership', 'Issues often persist because no one knows who owns the fix or the impacted downstream process.'),
    ('Manual triage', 'Quality reviews rely on spreadsheets and scattered follow-up that slow resolution and reduce visibility.'),
    ('Opaque AI decisions', 'Teams need recommendations that are explainable, evidence-based, and tied to actual business context.'),
]
for i, (title, body) in enumerate(items):
    left = Inches(0.8 + (i % 2) * 5.3)
    top = Inches(3.1 + (i // 2) * 1.8)
    add_card(slide, left, top, Inches(4.6), Inches(1.3), title, body)

# Slide 3
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = WHITE
add_header(slide, 'DART turns data quality management into a repeatable operating system.', 'what it does')
features = [
    ('Reconcile', 'Compare source and target datasets, calculate weighted match rates, and surface high-risk field mismatches by stream and report.'),
    ('Diagnose', 'Identify drivers, score impact, and connect issues to the people, reports, and workflows they affect.'),
    ('Act', 'Generate actions, assign ownership, trigger alerts, and keep remediation visible throughout the workflow.'),
]
for i, (title, body) in enumerate(features):
    left = Inches(0.9 + i * 4.0)
    add_card(slide, left, Inches(2.4), Inches(3.5), Inches(3.2), title, body)

# Slide 4
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = WHITE
add_header(slide, 'The platform brings together reconciliation, governance, and intelligence.', 'core features')
features = [
    ('Reconciliation engine', 'Loads workbooks, standardizes values, and compares expected versus actual records with impact scoring.'),
    ('Impact and mapping intelligence', 'Shows downstream field-to-report and owner connections so effort is prioritized by business impact.'),
    ('Governance center', 'Tracks controls, decisions, owners, actions, and templates in a single operating layer.'),
    ('AI-assisted analysis', 'Transforms complex patterns into plain-English summaries grounded in evidence and current context.'),
    ('Automation and alerts', 'Persists email alerts, baselines, and workbook watchers to catch issues before they escalate.'),
    ('Multi-domain coverage', 'Supports claims, Medicaid intelligence, and BYO datasets in one product experience.'),
]
for i, (title, body) in enumerate(features):
    left = Inches(0.8 + (i % 2) * 5.4)
    top = Inches(2.2 + (i // 2) * 1.8)
    add_card(slide, left, top, Inches(4.6), Inches(1.35), title, body)

# Slide 5
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = WHITE
add_header(slide, 'A simple operational flow: from intake to resolution.', 'how it works')
steps = [
    ('01 | Ingest', 'Pull data from workbooks and source systems into a governed review layer.'),
    ('02 | Standardize', 'Normalize fields, compare records, and detect mismatches with evidence.'),
    ('03 | Score', 'Apply weighted match rates, impact scores, and risk tiers to determine action priority.'),
    ('04 | Resolve', 'Turn findings into actions, owners, alerts, and remediation plans with traceable status updates.'),
]
for i, (heading, body) in enumerate(steps):
    y = Inches(2.2 + i * 1.45)
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), y, Inches(11.0), Inches(0.9))
    box.fill.solid(); box.fill.fore_color.rgb = OFFWHITE
    box.line.color.rgb = LINE
    add_textbox(slide, Inches(1.25), y + Inches(0.18), Inches(2.5), Inches(0.3), heading, 14, GOLD4, True)
    add_textbox(slide, Inches(3.9), y + Inches(0.18), Inches(7.8), Inches(0.4), body, 13, MUTED)

# Slide 6
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = WHITE
add_header(slide, 'The business outcomes are operational and measurable.', 'outcomes')
outcomes = [
    ('Faster issue triage', 'Reduce the time from detection to action and help teams focus on the highest-risk mismatches first.'),
    ('Better prioritization', 'Shift effort toward issues with the largest downstream impact rather than the loudest symptoms.'),
    ('More trusted data', 'Improve confidence in KPIs, claims files, and operating reports across stakeholders.'),
]
for i, (title, body) in enumerate(outcomes):
    left = Inches(0.9 + i * 4.0)
    add_card(slide, left, Inches(2.3), Inches(3.5), Inches(3.2), title, body)

# Slide 7
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = WHITE
add_header(slide, 'The product reduces effort across governance, remediation, and escalation.', 'efficiency & roi')
rows = [
    ('Reconciliation review', 'Manual workbook checks', 'Automated comparisons and scoring', 'High reduction in repetitive effort'),
    ('Issue triage', 'Ad hoc follow-up and spreadsheets', 'Structured issue queue and ownership', 'Faster escalation and accountability'),
    ('Reporting confidence', 'Delayed validation and low trust', 'Persistent dashboards and quality narrative', 'Fewer surprises and faster decisions'),
    ('Workflow automation', 'Manual email and status tracking', 'Alerts, baselines, and workbook watch', 'Lower operational overhead'),
]
header_row = ['Workstream', 'Before DART', 'With DART', 'Efficiency gain']
shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.1), Inches(11.7), Inches(3.8))
shape.fill.solid(); shape.fill.fore_color.rgb = OFFWHITE
shape.line.color.rgb = LINE
for i, h in enumerate(header_row):
    add_textbox(slide, Inches(1.0 + i * 2.9), Inches(2.35), Inches(2.6), Inches(0.45), h, 10, GOLD4, True)
for r, row in enumerate(rows):
    y = Inches(2.9 + r * 0.76)
    for c, val in enumerate(row):
        add_textbox(slide, Inches(1.0 + c * 2.9), y, Inches(2.6), Inches(0.55), val, 9, MUTED)

# Slide 8
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = WHITE
add_header(slide, 'The value is not just dashboards — it is operational control.', 'why this matters')
items = [
    ('For data teams', 'Reduce manual validation effort, standardize checks, and quickly trace root causes.'),
    ('For business owners', 'Know which issues materially affect claims performance, finance, or compliance outcomes.'),
    ('For leadership', 'Move from status reporting to a reliable quality narrative backed by evidence and progress.'),
    ('For governance', 'Turn issue management into a disciplined, accountable operating process.'),
]
for i, (title, body) in enumerate(items):
    left = Inches(0.8 + (i % 2) * 5.5)
    top = Inches(2.2 + (i // 2) * 1.8)
    add_card(slide, left, top, Inches(4.7), Inches(1.35), title, body)

# Slide 9
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = WHITE
add_header(slide, 'Built as a compact, practical application layer.', 'architecture')
arch = [
    ('Backend', 'FastAPI application logic with processing, routing, and operational APIs.'),
    ('Frontend', 'Embedded HTML/CSS/JS interface for a self-contained experience that is easy to run and demo.'),
    ('Data layer', 'Workbook and SQLite-backed tracking for reconciliation, issue monitoring, and BYO support.'),
]
for i, (title, body) in enumerate(arch):
    left = Inches(0.9 + i * 4.0)
    add_card(slide, left, Inches(2.4), Inches(3.5), Inches(3.1), title, body)

# Slide 10
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = WHITE
add_header(slide, 'DART is a data quality operating model, not just a reporting page.', 'closing message')
add_textbox(slide, Inches(0.8), Inches(2.4), Inches(10.7), Inches(1.2), 'It is a product designed for organizations that need to reconcile messy data, understand impact, assign accountability, and drive remediation in a disciplined, measurable way.', 18, MUTED)
chips = ['Actionable', 'Evidence-based', 'Operational', 'Scalable', 'Governed']
for i, chip in enumerate(chips):
    left = Inches(0.9 + i * 2.05)
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(4.3), Inches(1.6), Inches(0.55))
    shape.fill.solid(); shape.fill.fore_color.rgb = OFFWHITE
    shape.line.color.rgb = LINE
    add_textbox(slide, left + Inches(0.1), Inches(4.42), Inches(1.4), Inches(0.3), chip, 12, GOLD4, True, PP_ALIGN.CENTER)
add_textbox(slide, Inches(0.8), Inches(5.6), Inches(10.4), Inches(0.7), 'Repository centerpiece: app6.py', 16, GOLD4, True)

output_path = 'DART_Product_Deck.pptx'
prs.save(output_path)
print(f'Created {output_path}')
